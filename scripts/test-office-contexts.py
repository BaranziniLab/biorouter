#!/usr/bin/env python3
"""Execute the built-in office examples and verify real file round trips.

Run with the packages listed in docs/design/office-context-compatibility.md.
Outputs are QA artifacts; PNGs still require human/model visual inspection.
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader, PdfWriter


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    out = args.output.resolve()
    out.mkdir(parents=True, exist_ok=False)
    office = shutil.which("soffice") or shutil.which("libreoffice")
    poppler = shutil.which("pdftoppm")
    if not office or not poppler:
        raise RuntimeError("Require LibreOffice (soffice/libreoffice) and pdftoppm on PATH")
    root = Path(__file__).resolve().parents[1]
    for name in ("word", "powerpoint", "excel", "pdf"):
        skill = root / f"crates/biorouter/src/agents/builtin_skills/office-{name}/SKILL.md"
        examples = re.findall(r"```python\n(.*?)\n```", skill.read_text(), re.S)
        require(len(examples) == 1, f"Expected one runnable example in {skill}")
        builder = out / f"{name}.py"
        builder.write_text(examples[0])
        subprocess.run([sys.executable, str(builder)], cwd=out, check=True, timeout=60)

    files = out / "output"
    word = files / "report.docx"
    doc = Document(word)
    require(doc.paragraphs[0].text == "Project report", "Word title")
    require(doc.tables[0].cell(1, 1).text == "12", "Word table")
    doc.add_paragraph("Round-trip edit verified.")
    doc.save(files / "report-edited.docx")
    edited = Document(files / "report-edited.docx")
    require(edited.paragraphs[-1].text == "Round-trip edit verified.", "Word edit")
    require(edited.tables[0].cell(1, 1).text == "12", "Word edit preserved table")

    deck = files / "briefing.pptx"
    prs = Presentation(deck)
    require("evidence" in prs.slides[0].notes_slide.notes_text_frame.text, "Slide notes")
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Next steps"
    slide.placeholders[1].text = "Review the evidence and confirm the plan."
    prs.save(files / "briefing-edited.pptx")
    edited_deck = Presentation(files / "briefing-edited.pptx")
    require(len(edited_deck.slides) == 2, "Slide edit count")
    require("evidence" in edited_deck.slides[0].notes_slide.notes_text_frame.text,
            "Slide edit preserved notes")

    book = files / "budget.xlsx"
    wb = load_workbook(book)
    require(wb.active["D2"].value == "=B2*C2", "Workbook formula")
    require(wb.active.freeze_panes == "A2", "Workbook freeze panes")
    wb.active["B2"] = 4
    wb.save(files / "budget-edited.xlsx")

    pdf = files / "summary.pdf"
    require("Project summary" in PdfReader(pdf).pages[0].extract_text(), "PDF text")
    writer = PdfWriter()
    writer.append(pdf)
    writer.append(pdf)
    with (files / "summary-merged.pdf").open("wb") as stream:
        writer.write(stream)
    require(len(PdfReader(files / "summary-merged.pdf").pages) == 2, "PDF merge")

    def convert(source, kind, label):
        dest = out / label
        dest.mkdir()
        profile = out / f"profile-{label}"
        subprocess.run([
            office, f"-env:UserInstallation={profile.as_uri()}", "--headless",
            "--convert-to", kind, "--outdir", str(dest), str(source),
        ], check=True, timeout=120, capture_output=True, text=True)
        result = dest / f"{source.stem}.{kind}"
        require(result.is_file() and result.stat().st_size > 0, f"Conversion missing: {result}")
        return result

    calculated = convert(book, "xlsx", "calculated")
    changed = convert(files / "budget-edited.xlsx", "xlsx", "changed-calculated")
    for path, expected in [(calculated, 117.5), (changed, 130)]:
        values = load_workbook(path, data_only=True)
        formulas = load_workbook(path, data_only=False)
        require(values.active["D4"].value == expected, f"Recalculation expected {expected}")
        require(formulas.active["D4"].value.startswith("="), "Recalculation lost formula")
        for row in values.active:
            for cell in row:
                require(cell.data_type != "e", f"Formula error at {cell.coordinate}: {cell.value}")

    previews = [pdf, files / "summary-merged.pdf"]
    for index, source in enumerate([word, files / "report-edited.docx", deck,
                                    files / "briefing-edited.pptx", calculated, changed]):
        previews.append(convert(source, "pdf", f"render-{index}"))
    images = out / "pages"
    images.mkdir()
    for index, preview in enumerate(previews):
        require(len(PdfReader(preview).pages) > 0, f"Empty preview: {preview}")
        subprocess.run([poppler, "-scale-to", "1600", "-png", str(preview),
                        str(images / f"preview-{index}")], check=True, timeout=120,
                       capture_output=True)
    report = {
        "structural_checks": "passed",
        "word_edit_preserved_table": True,
        "powerpoint_edit_preserved_notes": True,
        "pdf_merge_pages": 2,
        "spreadsheet_totals": [117.5, 130],
        "calculation_engine": office,
        "rendered_images": [str(p) for p in sorted(images.glob("*.png"))],
        "visual_inspection": "pending: inspect each PNG",
        "native_microsoft_office": "not tested",
    }
    (out / "verification.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
